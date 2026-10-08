"""Create a DOCX version of the beginner student resume for testing."""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Name
name = doc.add_paragraph()
name_run = name.add_run("Priya Sharma")
name_run.font.size = Pt(24)
name_run.font.bold = True
name.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Contact
contact = doc.add_paragraph()
contact_run = contact.add_run("Email: priya.sharma@student.edu | Phone: +91 98765 43210 | LinkedIn: linkedin.com/in/priyasharma")
contact_run.font.size = Pt(10)
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()  # Spacing

# Education
edu_heading = doc.add_heading("Education", level=2)
edu = doc.add_paragraph()
edu.add_run("Bachelor of Technology in Computer Science\n").bold = True
edu.add_run("XYZ Institute of Technology | Expected Graduation: May 2028 | CGPA: 8.2/10")

# Technical Skills
skills_heading = doc.add_heading("Technical Skills", level=2)

skills_para = doc.add_paragraph()
skills_para.add_run("Programming Languages: ").bold = True
skills_para.add_run("Python, C, Java")
skills_para = doc.add_paragraph()
skills_para.add_run("Web Technologies: ").bold = True
skills_para.add_run("HTML, CSS, JavaScript (basics)")
skills_para = doc.add_paragraph()
skills_para.add_run("Tools & Platforms: ").bold = True
skills_para.add_run("Git, VS Code, PyCharm")
skills_para = doc.add_paragraph()
skills_para.add_run("Databases: ").bold = True
skills_para.add_run("MySQL (learning)")

# Projects
projects_heading = doc.add_heading("Projects", level=2)

project1 = doc.add_paragraph()
project1.add_run("Student Grade Calculator\n").bold = True
project1.add_run("Python | September 2024").italic = True
doc.add_paragraph("• Developed a command-line application to calculate student grades and GPA")
doc.add_paragraph("• Implemented input validation and error handling for robust data entry")
doc.add_paragraph("• Used file I/O to store and retrieve student records")

project2 = doc.add_paragraph()
project2.add_run("Personal Portfolio Website\n").bold = True
project2.add_run("HTML, CSS, JavaScript | October 2024").italic = True
doc.add_paragraph("• Created a responsive personal portfolio website showcasing academic projects")
doc.add_paragraph("• Implemented smooth scrolling navigation and interactive elements using JavaScript")
doc.add_paragraph("• Deployed the website using GitHub Pages")

# Coursework
coursework_heading = doc.add_heading("Coursework", level=2)
doc.add_paragraph("• Introduction to Programming (C)")
doc.add_paragraph("• Data Structures and Algorithms")
doc.add_paragraph("• Database Management Systems")
doc.add_paragraph("• Object-Oriented Programming with Java")
doc.add_paragraph("• Web Technologies")

# Certifications
cert_heading = doc.add_heading("Certifications", level=2)
cert = doc.add_paragraph()
cert.add_run("Python for Everybody Specialization\n").bold = True
cert.add_run("Coursera | August 2024")

# Achievements
achieve_heading = doc.add_heading("Achievements", level=2)
doc.add_paragraph("• First Prize in College Coding Competition (Fresher's Code Sprint 2024)")
doc.add_paragraph("• Active member of Computer Science Club")

# Save
doc.save("beginner-student-resume.docx")
print("Created beginner-student-resume.docx")
